"""Integration + browser checks against a RUNNING test app. Creates synthetic test cards."""
import hashlib
import io
import json
import os
import uuid
from pathlib import Path
import httpx
import openpyxl
from playwright.sync_api import sync_playwright

ROOT=Path(__file__).resolve().parents[1]
BASE=os.getenv('RESEARCH_URL','http://127.0.0.1:8000')
client=httpx.Client(base_url=BASE,trust_env=False,timeout=60)
schema=client.get('/api/schema').json()
assert client.get('/api/health').json()['status']=='ok'
code='ТЕСТ_MVP_001'
sections={s['name']:[] for s in schema}
sections['Пациенты']=[{'Код пациента':code,'Индексная дата':'2026-10-05','Дата рождения':'1985-02-14','Пол':'Мужской','Рост см':175,'Масса тела кг':75,'Окружность талии см':85,'ЧСС':72,'САД':120,'ДАД':80,'Статус курения':0,'СС риск записанный':1,'Группа здоровья записанная':'1'}]
sections['Лабораторные результаты']=[dict(zip([f['name'] for f in schema[1]['fields']],[code,145,4.7,6.1,240,43,90,30,330,8,5.0,4.5,2.3]))]
sections['Диагнозы']=[{'Код пациента':code,'Код диагноза':'ТЕСТОВЫЙ_КОД','Статус диагноза':'pre_existing'}]
sections['ЭКГ']=[{'Код пациента':code,'ЭКГ выполнена':1,'Дата ЭКГ':'2026-10-05','Исходный текст ЭКГ':'ТЕСТ. Вымышленная запись.'}]
sections['Скрининг']=[{'Код пациента':code,'Анализ кала показан':0,'Анализ кала выполнен':0,'Результат анализа кала':'NA','Маммография показана':0,'Маммография выполнена':0}]
recommendation=next(f for s in schema if s['name']=='Второй этап' for f in s['fields'] if f['name']=='Код рекомендации второго этапа')['options'][0]['value']
sections['Второй этап']=[{'Код пациента':code,'Код рекомендации второго этапа':recommendation,'Дата второго этапа':'2026-10-05','Исходный текст результата второго этапа':'ТЕСТ. Не медицинское заключение.'}]
sections['Текст рекомендаций врача']=[{'Код пациента':code,'Исходный текст рекомендаций врача':'ТЕСТ. Вымышленный пример рекомендации.'}]
sections['Рекомендации врача']=[{'Код пациента':code,'Код рекомендации':recommendation,'Рекомендовано врачом':1}]
exists=client.get('/api/patients/'+code).status_code==200
payload={'code':code,'mode':'edit' if exists else 'create','sections':sections}
r=client.post('/api/patients',json=payload); assert r.status_code==200,r.text
saved=client.get('/api/patients/'+code).json()
assert saved['Лабораторные результаты'][0]['Гемоглобин']==145
assert saved['Пациенты'][0]['Группа здоровья записанная']=='1'
assert saved['Скрининг'][0]['Дата анализа кала'] is None
assert len(saved['Рекомендации врача'])==1
payload['mode']='create'
assert client.post('/api/patients',json=payload).status_code==409
payload['mode']='edit'; sections['Пациенты'][0]['СС риск записанный']=5
assert client.post('/api/patients',json=payload).status_code==422
sections['Пациенты'][0]['СС риск записанный']=1
sections['Пациенты'][0]['Выдуманное поле']='x'
assert client.post('/api/patients',json=payload).status_code==422
del sections['Пациенты'][0]['Выдуманное поле']
assert client.get('/api/patients/'+code).json()['Пациенты'][0]['СС риск записанный']==1

errors=[]
with sync_playwright() as p:
 browser=p.chromium.launch(executable_path=os.getenv('CHROMIUM_PATH','/usr/bin/chromium'),headless=True,args=['--no-sandbox'])
 page=browser.new_page(viewport={'width':1440,'height':1000})
 page.on('pageerror',lambda e:errors.append(str(e)))
 page.goto(BASE)
 page.get_by_role('button',name=code,exact=False).wait_for()
 page.get_by_role('button',name=code,exact=False).click()
 page.wait_for_function("document.querySelector('#patient-code').value === 'ТЕСТ_MVP_001'")
 assert page.get_by_label('СС риск записанный').input_value()=='1'
 page.get_by_label('СС риск записанный').select_option('4')
 page.get_by_label('Группа здоровья записанная').select_option('3а')
 page.get_by_role('button',name='Сохранить в таблицу',exact=True).click()
 page.wait_for_function("document.querySelector('#status').textContent.startsWith('Сохранено в рабочую')",timeout=60000)
 assert client.get('/api/patients/'+code).json()['Пациенты'][0]['Группа здоровья записанная']=='3а'
 page.get_by_role('tab',name='Лабораторные результаты',exact=True).click()
 page.evaluate('window.scrollTo(0,0)')
 page.wait_for_timeout(180)
 page.screenshot(path=str(ROOT/'docs/desktop.png'),full_page=True)
 page.set_viewport_size({'width':390,'height':844})
 for name in ['Пациенты','Лабораторные результаты','Диагнозы','ЭКГ','Скрининг','Второй этап','Рекомендации врача']:
  page.get_by_role('tab',name=name,exact=True).click()
  assert page.evaluate('document.documentElement.scrollWidth<=window.innerWidth'),name
 page.get_by_role('tab',name='Пациенты',exact=True).click()
 page.wait_for_timeout(180)
 page.screenshot(path=str(ROOT/'docs/mobile.png'),full_page=True)
 # Hidden invalid required fields must reveal the correct tab before focus.
 page.get_by_role('button',name='Новая запись',exact=True).click()
 form_code='ТЕСТ_ФОРМА_'+uuid.uuid4().hex[:6].upper()
 page.get_by_label('Код пациента',exact=False).fill(form_code)
 page.get_by_label('Пол',exact=True).select_option('Женский')
 page.get_by_label('Статус курения',exact=True).select_option('9')
 page.get_by_role('tab',name='Диагнозы',exact=True).click()
 page.get_by_role('button',name='Добавить строку',exact=True).filter(visible=True).click()
 page.get_by_role('tab',name='Пациенты',exact=True).click()
 page.get_by_role('button',name='Сохранить в таблицу',exact=True).click()
 assert not page.locator('.panel[data-name="Диагнозы"]').evaluate('(node)=>node.hidden')
 page.get_by_label('Код диагноза',exact=False).fill('ТЕСТ_КОД_2')
 page.get_by_label('Статус диагноза',exact=False).select_option('new')
 page.get_by_role('button',name='Сохранить в таблицу',exact=True).click()
 page.wait_for_function("document.querySelector('#status').textContent.startsWith('Сохранено в рабочую')",timeout=60000)
 assert client.get('/api/patients/'+form_code).status_code==200
 with page.expect_download() as download:
  page.get_by_role('link',name='Скачать Excel',exact=True).click()
 assert download.value.failure() is None
 browser.close()
assert not errors,errors
content=client.get('/api/export').content
w=openpyxl.load_workbook(io.BytesIO(content))
for name in sections:
 sheet=w[name]
 row=next(r for r in sheet.iter_rows(min_row=2) if r[0].value==code)
 assert row[0].value==code
 if name=='Пациенты':
  assert row[2].value.strftime('%Y-%m-%d')=='1985-02-14'
  assert row[11].value==4 and row[12].value=='3а'
 if name=='Рекомендации врача': assert row[3].data_type=='f'
source=ROOT/'docs/workbook-structure.json'
expected=json.loads(source.read_text())['source_sha256']
original=Path(os.environ['RESEARCH_SOURCE']) if 'RESEARCH_SOURCE' in os.environ else next((ROOT.parent/'attachments').glob('*/*.xlsx'),None)
formula_count=None
source_verified=False
if original and original.exists():
 assert hashlib.sha256(original.read_bytes()).hexdigest()==expected
 original_book=openpyxl.load_workbook(original)
 formula_count=0
 for sheet in original_book:
  for cell in sheet._cells.values():
   if cell.data_type=='f':
    assert w[sheet.title][cell.coordinate].value==cell.value,(sheet.title,cell.coordinate)
    formula_count+=1
 assert formula_count>0
 source_verified=True
result={'api':'passed','create':'passed','edit':'passed','xlsx_mapping':'passed','reject_invalid_values':'passed','reject_unknown_fields':'passed','duplicate_code':'passed','desktop':'1440 × 1000','mobile':'390 × 844; all seven sections, no horizontal overflow','download':'passed','browser_errors':errors,'preserved_formulas':formula_count,'source_verified':source_verified,'original_sha256':expected}
(ROOT/'docs/test-results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False,indent=2))
