"""Read the source workbook without modifying it; export structure only."""
import hashlib
import json
import sys
from pathlib import Path
import openpyxl

source = Path(sys.argv[1])
output = Path(sys.argv[2])
workbook = openpyxl.load_workbook(source, data_only=False)
schema = {
    'source_name': source.name,
    'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
    'named_ranges': {key: value.attr_text for key, value in workbook.defined_names.items()},
    'sheets': []
}
for sheet in workbook:
    schema['sheets'].append({
        'name': sheet.title,
        'headers': [cell.value for cell in sheet[1]],
        'rows': sheet.max_row,
        'tables': {name: sheet.tables[name].ref for name in sheet.tables},
        'validations': [{
            'range': str(rule.sqref), 'type': rule.type,
            'operator': rule.operator, 'formula1': rule.formula1,
            'formula2': rule.formula2, 'allow_blank': rule.allow_blank
        } for rule in sheet.data_validations.dataValidation]
    })
output.write_text(json.dumps(schema, ensure_ascii=False, indent=2), encoding='utf-8')
print('Exported', len(schema['sheets']), 'sheets; source SHA256:', schema['source_sha256'])
