import json
from pathlib import Path
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
working=ROOT/'data/01_База_исследования_рабочая.xlsx'
w=openpyxl.load_workbook(working if working.exists() else ROOT/'templates/research-test.xlsx')
sheets=['Пациенты','Лабораторные результаты','Диагнозы','ЭКГ','Скрининг','Второй этап','Текст рекомендаций врача','Рекомендации врача']
units={r[0]:r[1] for r in w['Единицы измерения'].iter_rows(min_row=2,values_only=True) if r[0]}
dictionary={(r[0],r[1]):r for r in w['Словарь данных'].iter_rows(min_row=2,values_only=True)}
def choices(name):
 if name=='Список_рекомендаций':
  return [{'value':r[0],'label':str(r[0])+' · '+str(r[1])} for r in w['Справочник рекомендаций'].iter_rows(min_row=2,values_only=True) if r[0] and r[4]=='УТВЕРЖДЕНО']
 if name=='Список_пациентов': return None
 n=w.defined_names[name]
 result=[]
 for sheet,region in n.destinations:
  for row in w[sheet][region]:
   c=row[0]
   if c.value is not None:
    label=str(c.value)
    if name in ['Список_курения','Список_решений']: label+=' · '+str(w[sheet].cell(c.row,c.column+1).value)
    result.append({'value':c.value,'label':label})
 return result
schema=[]
for name in sheets:
 s=w[name]; fields=[]
 for c in s[1]:
  if name=='Рекомендации врача' and c.column==4: continue # formula, never editable
  field={'name':c.value,'column':c.column,'type':'text','required':c.column==1,'unit':units.get(c.value)}
  info=dictionary.get((name,c.value))
  if info and info[5]=='Дата Excel': field['type']='date'
  if field['unit']: field['type']='number'
  if c.value and ('Исходный текст' in c.value): field['type']='textarea'
  for rule in s.data_validations.dataValidation:
   if s.cell(2,c.column).coordinate in rule.sqref:
    if rule.type=='list':
     options=choices(rule.formula1)
     if options is not None: field.update(type='select',options=options)
    elif rule.type=='decimal':
     field['type']='number'
     if rule.operator=='greaterThan': field['exclusive_min']=float(rule.formula1)
  overrides={'СС риск записанный':[1,2,3,4],'Группа здоровья записанная':['1','2','3а','3б'],'BI-RADS справа':[0,1,2,3,4,5,6],'BI-RADS слева':[0,1,2,3,4,5,6]}
  if c.value in overrides: field.update(type='select',options=[{'value':v,'label':str(v)} for v in overrides[c.value]],source='Уточнение пользователя')
  if c.value=='СС риск записанный':
   labels={1:'низкий',2:'относительный',3:'высокий',4:'очень высокий'}
   field['options']=[{'value':v,'label':str(v)+' · '+labels[v]} for v in overrides[c.value]]
  if name in ['Диагнозы','Второй этап','Рекомендации врача']: field['required']=c.column<= (3 if name in ['Диагнозы','Второй этап'] else 2)
  fields.append(field)
 schema.append({'name':name,'many':name in ['Диагнозы','Второй этап','Рекомендации врача'],'fields':fields})
(ROOT/'app/schema.json').write_text(json.dumps(schema,ensure_ascii=False,indent=2))
print('Schema:',len(schema),'input sheets')
